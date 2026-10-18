# Copyright (C) 2026 annsshadow
# SPDX-License-Identifier: AGPL-3.0-or-later

"""L287 / A389：电子表格公式注入（CWE-1236）——四条交付边按各自格式的原生手段挡住公式

数据集里的 `instruction` / `output` 是**任何能写数据集的人**都够得着的载荷面。改前实测
（L286 立账那轮，仓外探针在 L287 复量）：

- CSV / TSV 三条写边（`export.py` 原生支、`export_enhanced.py` 的 csv 与 tsv、
  `converter.py` 的 csv 与 tsv）把 `=1+1` 与 `+2` 形状的字符串**原样成文**，字节里就是那行
  文本，Excel / LibreOffice 打开即当公式解释；
- xlsx 写边（`excel_write.py`，本仓 Excel 落盘的唯一产品落点）更硬：openpyxl 按值自己定
  类型，前导 `=` 被判成公式，产物 XML 里是 `<f>1+1</f>`，openpyxl 读回 `data_type == "f"`
  且值为那个表达式、pandas 读回是 `NaN` —— 交付物里那一格的数据本身也丢了。

处置按用户 2026-10-17 对 A389 的拍板①「各按格式原生转义」：

- CSV / TSV 没有类型位，只有前导单引号这一条既有约定 ⇒ `validation.sheet_text_row` 给会被
  执行的字符串前面加一个 `'`（Excel / LibreOffice / Numbers 认它为「文本」标记）。
  **这一改写进交付字节**，读回那份 CSV 拿到 `'=1+1`；本守卫第 5 组把这条后果钉住，
  防止下一轮把它当回归改回去。
- xlsx 有真正的类型位 ⇒ `excel_write.text_only_cell` 把 `data_type` 显式设成 `s`，
  **一个字都不必加**，格子里的字面串与进入时相同（唯一例外是格内的回车：xlsx 存的就是换行，
  openpyxl 读回给 `\n`，两条路同形、非本轮引入，第 3 组按这笔翻译记账）。
- 判据（哪些前导会被执行）只有一个产地 `validation.SHEET_FORMULA_LEADERS`，四条边共引，
  谁都不许手抄字符集（第 6、7 两组量这个）。

不在本轮面里的两格，留在账上而不是顺手做：给 **CSV 表头**（列名）加前缀会改变读回侧用来配行
的键，一处转义要两处同步，量到之后另立 A391 待拍；`json` / `jsonl` 两档没有公式这一族解释器，
第 4 组把它们当**负面对照**钉住。

尺子面：

1. 三条 CSV/TSV 边逐边盘面事实：六档前导各自加前缀，普通串与非字符串值一字不动。
2. TSV 两条边与 CSV 同一条判据，分隔符不改变判决。
3. xlsx 边：产物 XML 零 `<f>`、读回字面串与类型位（格内回车按 xlsx 自己的 `\r → \n` 记）、
   表头那格同样不产公式。
4. 负面对照：json / jsonl 一字未动 —— 判据不许顺手扩到容器之外。
5. 读回后果钉死：CSV 那份产物读回来带前导单引号（拍板接受的面，不是疏漏）。
6. 闭族普查，三条各有牙齿臂：产品树里每一处 `writerows` 都过 `sheet_text_row`、每一处
   `ws.append` 都过 `text_only_cell`，且 xlsx 容器的出口只许住 `augmentor/excel_write.py`
   （`to_excel` / `ExcelWriter` / `Workbook()` 三形现量各只在那只文件里出现）。
7. 判据住一处，分两档（两档各有牙齿臂）：四条写边文件里**一档前导字面量都不许出现**
   （它们只许共引权威），产品树其余部分只禁「一个集合里抄了 ≥2 档可打印前导」那一形状。
   窄的那一档不是放宽：制表与回车是解析类代码的空白常客（`json_extract.py:104` 的
   `in ' \\t\\r\\n'`、`validation.py` 的 glob 字符组），把它们计入会把无关写法报成缺陷；
   结构性保证在第 6 组（每条写边必须过权威），这一组只防「把同一把常数抄第二份」。
8. 牙齿：把 `sheet_text_row` 换成恒等 ⇒ 裸 `=1+1` 当场重现；把 `is_sheet_formula_text`
   判成恒假 ⇒ xlsx 又长出 `<f>`。
"""

from __future__ import annotations

import ast
import csv
import json
import sys
import zipfile
from pathlib import Path
from typing import Any, Dict, List

import pytest

REPO = Path(__file__).resolve().parents[2]
if str(REPO) not in sys.path:
    sys.path.insert(0, str(REPO))

from augmentor import excel_write, validation  # noqa: E402
from augmentor.converter import convert_file  # noqa: E402
from augmentor.export import Exporter  # noqa: E402
from augmentor.export_enhanced import EnhancedExporter, ExportFormat, ExportOptions  # noqa: E402
from augmentor.validation import SHEET_FORMULA_LEADERS  # noqa: E402

#: 会被电子表格执行的前导，逐档点名（第 1、7 两组共用同一份载荷）
ESCAPED = ["=1+1", "+2", "-3", "@SUM(1)", "\tx", "\r-enter"]

PAYLOADS: List[Dict[str, str]] = [
    {"instruction": "=1+1", "output": "ok"},
    {"instruction": "+2", "output": "-3"},
    {"instruction": "@SUM(1)", "output": "\tx"},
    {"instruction": "\r-enter", "output": "普通文本"},
]

#: 四条边都够得着的落点族：禁令面就是这一族，不是全语料（循环纪律第 6 节）
PRODUCT_ROOTS = ("augmentor", "api")

#: 交付字节从这里出去，判据一律共引 `validation`，自己文件里连一档前导字面量都不该有
WRITE_EDGES = (
    "augmentor/export.py",
    "augmentor/export_enhanced.py",
    "augmentor/converter.py",
    "augmentor/excel_write.py",
)

#: 第 7 组在全产品树那一档只数这四档：制表与回车是解析类代码的空白常客，
#: 把它们算进「抄了字符集」的判据会被 `json_extract.py:104` 那种无关写法填满。
CENSUS_LEADERS = ("=", "+", "-", "@")


def csv_rows(path: Path, delimiter: str = ",") -> List[List[str]]:
    """读回一份 CSV/TSV 产物（表头 + 逐行），与写边同一种分隔符"""
    with open(path, encoding="utf-8", newline="") as handle:
        return list(csv.reader(handle, delimiter=delimiter))


def flat_cells(rows: List[List[str]]) -> set:
    """表体格子（跳过表头那一行）"""
    return {cell for row in rows[1:] for cell in row}


def sheet_xml(path: Path) -> str:
    """xlsx 容器里首个工作表的 XML（判 `<f>` 只能看字节，读回侧会替它说谎）"""
    with zipfile.ZipFile(path) as archive:
        names = sorted(n for n in archive.namelist() if n.startswith("xl/worksheets/sheet"))
        return archive.read(names[0]).decode("utf-8")


def native_csv(items: List[Dict], path: Path, fmt: str = "csv") -> None:
    Exporter().export(items, str(path), fmt)


def enhanced_csv(items: List[Dict], path: Path, fmt: str = "csv") -> None:
    EnhancedExporter().export(items, str(path),
                              ExportOptions(format=EXPORT_FORMATS[fmt]))


def converted(items: List[Dict], path: Path, fmt: str = "csv") -> None:
    """转换器写边：`convert_file` 是唯一会落盘的入口（`DatasetConverter().convert` 只回数据）"""
    src = path.parent / "src_for_write.json"
    src.write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
    convert_file(str(src), str(path), target_format=fmt, source_format="json")


#: 三条边的格式名各自住在不同表示上（原生档收字符串、委托档收枚举），这里只做名字翻译
EXPORT_FORMATS = {"csv": ExportFormat.CSV, "tsv": ExportFormat.TSV}

WRITERS = {"native": native_csv, "enhanced": enhanced_csv, "converter": converted}


class TestLeaderSetIsOneAuthority:
    def test_the_leader_set_is_pinned(self):
        """判据本体逐档钉死：加一档要写进账本，减一档会让某条边重新交出公式"""
        assert SHEET_FORMULA_LEADERS == ("=", "+", "-", "@", "\t", "\r")
        assert ESCAPED and all(p[:1] in SHEET_FORMULA_LEADERS for p in ESCAPED)


class TestCsvEdgesEscapeFormulaLead:
    @pytest.mark.parametrize("edge", sorted(WRITERS), ids=str)
    def test_every_leader_gets_the_apostrophe_on_every_edge(self, tmp_path, edge):
        """三条 CSV 边逐边：六档前导全部带前导单引号，普通串一字不动"""
        out = tmp_path / f"{edge}.csv"
        WRITERS[edge](PAYLOADS, out, "csv")
        flat = flat_cells(csv_rows(out))
        for payload in ESCAPED:
            assert "'" + payload in flat, f"{edge} 边没有转义 {payload!r}：{sorted(flat)}"
        assert "ok" in flat and "普通文本" in flat, f"{edge} 边把不该动的值也改了：{sorted(flat)}"
        assert not any(cell.startswith("=") for cell in flat if cell), (
            f"{edge} 边仍然交得出以 = 开头的格子：{sorted(flat)}"
        )

    @pytest.mark.parametrize("edge", ["enhanced", "converter"], ids=str)
    def test_tsv_says_the_same_as_csv(self, tmp_path, edge):
        """TSV 两条边（委托面与转换器）与 CSV 同一条判据：分隔符不改变判决

        原生 `export.py` 没有 tsv 分支（`NATIVE_FORMATS` 里就没有这一档），所以这一组只点两条
        真存在的边 —— 写成三条会当场红在「格式不存在」上，那是另一格账（A343 族）。
        """
        out = tmp_path / f"{edge}.tsv"
        WRITERS[edge](PAYLOADS, out, "tsv")
        flat = flat_cells(csv_rows(out, "\t"))
        assert "'" + "=1+1" in flat and "'" + "\tx" in flat, sorted(flat)
        assert not any(cell.startswith("=") for cell in flat if cell), sorted(flat)

    def test_non_string_values_are_untouched(self, tmp_path):
        """数字、布尔、None 不进判据：它们到了表格里就是原生类型，不会长成公式"""
        out = tmp_path / "scalars.csv"
        enhanced_csv([{"instruction": 42, "output": True, "flag": None, "ratio": -1.5}], out)
        rows = csv_rows(out)
        assert rows[0] == ["instruction", "output", "flag", "ratio"], rows
        assert rows[1] == ["42", "True", "", "-1.5"], rows


class TestExcelEdgeKeepsTheLiteralAndForcesType:
    """xlsx 有类型位，所以一个字都不必加 —— 与 CSV 那三条边的手段不同、判据同一把"""

    requires_openpyxl = pytest.mark.skipif(
        not excel_write.HAS_OPENPYXL, reason="写 xlsx 需要 openpyxl")

    @requires_openpyxl
    def test_no_formula_element_and_the_text_survives_verbatim(self, tmp_path):
        out = tmp_path / "escaped.xlsx"
        excel_write.export_table_to_excel(PAYLOADS, str(out), columns=["instruction", "output"])
        xml = sheet_xml(out)
        assert "<f>" not in xml and "<f " not in xml, xml[:400]

        from openpyxl import load_workbook

        book = load_workbook(out)
        try:
            got = [(c.value, c.data_type) for c in book.active["A"]]
        finally:
            book.close()
        assert got[0] == ("instruction", "s"), got
        # 只有一处翻译不是本轮引入的：单元格里的回车在 xlsx 里本来就是「换行」，openpyxl
        # 读回来给 `\n`。仓外探针量过两条路同形（裸写 `ws.append` 与强制文本类型都回
        # `\n-enter`、都是 `s`），所以这里按 `\r → \n` 记账，而不是假装逐字相同。
        assert [v for v, _ in got[1:]] == [
            p["instruction"].replace("\r", "\n") for p in PAYLOADS
        ], got
        assert all(dtype == "s" for _, dtype in got), got

    @requires_openpyxl
    def test_the_header_cell_is_forced_too(self, tmp_path):
        """列名本身可以来自数据集：一个 `=cmd` 的键在改前也是一枚公式"""
        out = tmp_path / "forged_header.xlsx"
        excel_write.export_table_to_excel([{ "=cmd": "v" }], str(out), columns=["=cmd"])
        assert "<f>" not in sheet_xml(out)

        from openpyxl import load_workbook

        book = load_workbook(out)
        try:
            head = book.active["A1"]
            assert (head.value, head.data_type) == ("=cmd", "s"), head.value
        finally:
            book.close()


class TestJsonFacesAreNotInScope:
    @pytest.mark.parametrize("fmt", ["json", "jsonl"], ids=str)
    def test_container_without_a_formula_interpreter_is_untouched(self, tmp_path, fmt):
        """没有公式解释器的那两档一字未动 —— 判据不许顺手扩到容器之外"""
        out = tmp_path / f"face.{fmt}"
        converted(PAYLOADS, out, fmt)
        text = out.read_text(encoding="utf-8")
        assert '"=1+1"' in text, text[:200]
        assert "'=1+1" not in text, text[:200]


class TestTheRoundTripConsequenceIsBooked:
    def test_the_apostrophe_comes_back_out_of_the_csv(self, tmp_path):
        """拍板① 的已知后果：CSV 读回是 `'=1+1`（交付字节里就带着那个标记）

        这一支不是「验它对」，是把**代价**钉在账上：下一轮若有人觉得「读回来多个引号算回归」
        而把判据摘掉，这一支与第 1 组会同时红，改动就必须在账本里重拍一次。
        """
        out = tmp_path / "rt.csv"
        enhanced_csv([{"instruction": "=1+1", "output": "ok"}], out)
        assert csv_rows(out)[1] == ["'=1+1", "ok"], csv_rows(out)


class ProductScan:
    """产品树的调用点扫描（判据钉在「族」上，不是全语料正则）"""

    @staticmethod
    def product_files() -> List[Path]:
        return sorted(p for root in PRODUCT_ROOTS for p in (REPO / root).rglob("*.py"))

    @staticmethod
    def calls_in(tree: ast.AST, name: str, label: str) -> List[tuple]:
        """返回 `(文件:行号, 整条调用的字面写法)` —— 写法里带着实参，权威有没有被调用看得见"""
        hits = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
                if node.func.attr != name or not node.args:
                    continue
                hits.append((f"{label}:{node.lineno}", ast.unparse(node)))
        return hits

    @classmethod
    def calls(cls, name: str) -> List[tuple]:
        return [
            hit
            for path in cls.product_files()
            for hit in cls.calls_in(
                ast.parse(path.read_text(encoding="utf-8")),
                name,
                path.relative_to(REPO).as_posix(),
            )
        ]

    @staticmethod
    def hand_copied_in(tree: ast.AST, label: str, strict: bool = False) -> List[str]:
        """手抄前导字符集的两种形状：`x.startswith(字面量)` 与 `x in (字面量组)`

        数的是**这一整个字面量集合里出现了几档前导字符**，不是逐个元素看：
        - `strict`（四条写边）：出现任意一档就报——那些文件只许共引权威，连单档都不该有；
        - 非 strict（产品树其余部分）：要 ≥2 档「可打印前导」（`CENSUS_LEADERS`）才报。
          制表与回车在解析类代码里本来就是空白档常客（`json_extract.py:104` 的
          `text[j] in ' \\t\\r\\n'` 就是），把它们算进这一档等于把尺子指向禁令面之外；
          真正的结构性保证是第 6 组（每条写边必须过权威），这一组只防「把常数抄第二份」。
        """
        hits = []
        pool = set(SHEET_FORMULA_LEADERS) if strict else set(CENSUS_LEADERS)
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
                if node.func.attr != "startswith" or not node.args:
                    continue
                args = node.args
            elif isinstance(node, ast.Compare) and isinstance(node.ops[0], ast.In):
                args = [node.comparators[0]]
            else:
                continue
            values = args[0].elts if isinstance(args[0], (ast.Tuple, ast.List, ast.Set)) else args
            copied = set()
            for value in values:
                if isinstance(value, ast.Constant) and isinstance(value.value, str):
                    copied |= {ch for ch in set(value.value) if ch in pool}
            if len(copied) >= (1 if strict else 2):
                hits.append(f"{label}:{node.lineno} {ast.unparse(node)[:60]}")
        return hits

    @staticmethod
    def container_writers_in(tree: ast.AST, label: str) -> List[tuple]:
        """「把 xlsx 容器写出去」的三种形状：`df.to_excel` / `ExcelWriter` / `Workbook()`

        前两条普查只盯 `writerows` 与 `ws.append`；一条新写的 `frame.to_excel(path)` 或自己
        `Workbook()` 的边会同时绕过它们（那两条量的是行写入形状，不是容器出口）。
        """
        hits = []
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            attr = getattr(node.func, "attr", None) or getattr(node.func, "id", None)
            if attr in ("to_excel", "ExcelWriter", "Workbook"):
                hits.append((f"{label}:{node.lineno}", ast.unparse(node)[:70]))
        return hits

    @classmethod
    def excel_writers(cls) -> List[tuple]:
        return [
            hit
            for path in cls.product_files()
            for hit in cls.container_writers_in(
                ast.parse(path.read_text(encoding="utf-8")),
                path.relative_to(REPO).as_posix(),
            )
        ]

    @classmethod
    def hand_copied(cls, paths=None, strict: bool = False) -> List[str]:
        return [
            where
            for path in (paths if paths is not None else cls.product_files())
            for where in cls.hand_copied_in(
                ast.parse(path.read_text(encoding="utf-8")),
                path.relative_to(REPO).as_posix(),
                strict,
            )
        ]


class TestTheFamilyIsClosed:
    def test_every_writerows_row_source_passes_the_authority(self):
        """这一族里每一处 `writerows` 都必须过权威：新增一条边没过，这里点名它"""
        sites = ProductScan.calls("writerows")
        assert sites, "扫描器一支都没读到 ⇒ 这是数不到，不是干净"
        for where, source in sites:
            assert "sheet_text_row" in source, f"{where} 的 writerows 没过权威：{source}"

    def test_every_excel_append_passes_the_authority(self):
        sites = [
            (where, source)
            for where, source in ProductScan.calls("append")
            if source.startswith("ws.append(")
        ]
        assert len(sites) == 2, [where for where, _ in sites]
        for where, source in sites:
            assert "text_only_cell" in source, f"{where} 的 ws.append 没过权威：{source}"

    def test_the_scanner_names_a_fresh_unescaped_writerows(self):
        """扫描器的牙齿：合成一段「新写的、没过权威的 writerows」，必须被它点名"""
        fake = "def f(w, rows):\n    w.writerows(rows)\n"
        hits = ProductScan.calls_in(ast.parse(fake), "writerows", "synthetic.py")
        assert hits == [("synthetic.py:2", "w.writerows(rows)")], hits

    def test_the_scanner_passes_the_escaped_form(self):
        fake = "def f(w, rows):\n    w.writerows([sheet_text_row(r) for r in rows])\n"
        hits = ProductScan.calls_in(ast.parse(fake), "writerows", "synthetic.py")
        assert len(hits) == 1 and "sheet_text_row" in hits[0][1], hits

    def test_the_hand_copy_scanner_has_teeth(self):
        """两档普查各有牙齿：≥2 档的形状两档都点名，单档只在写边那档（strict）点名"""
        forged = (
            "def f(v):\n"
            "    if v.startswith('='):\n"
            "        return \"'\" + v\n"
            "    return v\n"
            "\n"
            "def g(v):\n"
            "    return v[:1] in ('=', '+')\n"
        )
        loose = ProductScan.hand_copied_in(ast.parse(forged), "synthetic.py")
        assert [h.split(" ", 1)[1] for h in loose] == ["v[:1] in ('=', '+')"], loose
        strict = ProductScan.hand_copied_in(ast.parse(forged), "synthetic.py", True)
        assert sorted(h.split(" ", 1)[1] for h in strict) == [
            "v.startswith('=')",
            "v[:1] in ('=', '+')",
        ], strict

        authority = "def f(v):\n    return v[:1] in SHEET_FORMULA_LEADERS\n"
        assert ProductScan.hand_copied_in(ast.parse(authority), "synthetic.py") == []
        assert ProductScan.hand_copied_in(ast.parse(authority), "synthetic.py", True) == []

    def test_the_xlsx_container_has_one_landing_point(self):
        """Excel 落盘在本仓库只许住 `augmentor/excel_write.py`：绕开那条边的新写手当场点名"""
        funnel = ProductScan.excel_writers()
        assert funnel, "扫描器一支都没读到 ⇒ 这是数不到，不是干净"
        files = sorted({where.split(":")[0] for where, _ in funnel})
        assert files == ["augmentor/excel_write.py"], funnel

    def test_the_container_scanner_has_teeth(self):
        """上面那条的牙齿：合成一份 `frame.to_excel(path)` 必须被同一个扫描器点名"""
        forged = "def f(frame, path):\n    frame.to_excel(path, index=False)\n"
        hits = ProductScan.container_writers_in(ast.parse(forged), "synthetic.py")
        assert hits == [("synthetic.py:2", "frame.to_excel(path, index=False)")], hits

    def test_no_edge_hand_copies_the_leader_set(self):
        """判据住一处：产品树里不许出现抄了两档以上的字符集"""
        assert ProductScan.hand_copied() == []

    def test_the_four_write_edges_carry_no_leader_literal_at_all(self):
        """四条写边只许共引权威：连单档前导字面量都不该在自己文件里出现"""
        assert WRITE_EDGES and len(WRITE_EDGES) == 4, WRITE_EDGES
        paths = [REPO / rel for rel in WRITE_EDGES]
        for rel in WRITE_EDGES:
            assert (REPO / rel).is_file(), rel
        assert ProductScan.hand_copied(paths, strict=True) == []


class TestTheRulerHasTeeth:
    def test_identity_row_helper_puts_the_bare_formula_back(self, tmp_path, monkeypatch):
        """把 `sheet_text_row` 换成恒等 ⇒ CSV 里裸 `=1+1` 当场重现，证明第 1 组有牙齿

        写边用的是**函数内导入**（本仓既有口径），所以只改 `validation` 上那一个名字就够；
        如果哪天它改成模块级导入，这一支会不再命中 ⇒ 那时要补的是替身的位置，不是判据。
        """
        monkeypatch.setattr(validation, "sheet_text_row", lambda row: dict(row))
        out = tmp_path / "teeth.csv"
        EnhancedExporter().export([{"instruction": "=1+1"}], str(out),
                                  ExportOptions(format=ExportFormat.CSV))
        assert csv_rows(out)[1] == ["=1+1"], "摘掉权威仍然转义 ⇒ 第 1 组是恒真的"

    def test_a_blind_judgement_lets_the_formula_element_back_in(self, tmp_path, monkeypatch):
        """把 `is_sheet_formula_text` 判成恒假 ⇒ xlsx 又长出 `<f>`，证明第 3 组有牙齿"""
        if not excel_write.HAS_OPENPYXL:
            pytest.skip("本机没有 openpyxl，这一档无从注入")
        monkeypatch.setattr(excel_write, "is_sheet_formula_text", lambda value: False)
        out = tmp_path / "teeth.xlsx"
        excel_write.export_table_to_excel([{"a": "=1+1"}], str(out), columns=["a"])
        assert "<f>" in sheet_xml(out), "注入之后仍然不产公式 ⇒ 第 3 组是恒真的"
